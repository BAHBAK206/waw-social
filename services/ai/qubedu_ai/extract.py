"""Text extraction from PDF, DOCX, PPTX, XLSX, CSV and plain text, chunked by heading/page."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass

MAX_CHARS = 400_000
CHUNK = 900


class ExtractError(ValueError):
    pass


@dataclass
class Piece:
    heading: str | None
    page: int | None
    text: str


def _chunk(pieces: list[Piece]) -> list[dict]:
    out: list[dict] = []
    for p in pieces:
        text = re.sub(r"[ \t]+", " ", p.text).strip()
        if not text:
            continue
        paras = [x.strip() for x in re.split(r"\n{2,}|\n(?=[-•●▪])", text) if x.strip()]
        buf = ""
        for para in paras:
            if buf and len(buf) + len(para) + 2 > CHUNK:
                out.append({"heading": p.heading, "page": p.page, "text": buf})
                buf = ""
            if len(para) > CHUNK * 2:
                for i in range(0, len(para), CHUNK):
                    out.append({"heading": p.heading, "page": p.page, "text": para[i : i + CHUNK]})
                continue
            buf = f"{buf}\n\n{para}" if buf else para
        if buf:
            out.append({"heading": p.heading, "page": p.page, "text": buf})
    return out[:2000]


def _pdf(data: bytes) -> list[Piece]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractError("encrypted_pdf")
        pieces = []
        for i, page in enumerate(reader.pages[:500], start=1):
            text = page.extract_text() or ""
            if text.strip():
                pieces.append(Piece(None, i, text))
    except PdfReadError as e:
        raise ExtractError("invalid_pdf") from e
    if not pieces:
        raise ExtractError("pdf_has_no_text")
    return pieces


def _docx(data: bytes) -> list[Piece]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    pieces: list[Piece] = []
    heading: str | None = None
    buf: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "").lower()
        if style.startswith("heading") or style.startswith("title"):
            if buf:
                pieces.append(Piece(heading, None, "\n\n".join(buf)))
                buf = []
            heading = text[:200]
        else:
            buf.append(text)
    for table in doc.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        buf.append("\n".join(r for r in rows if r.strip(" |")))
    if buf:
        pieces.append(Piece(heading, None, "\n\n".join(buf)))
    return pieces


def _pptx(data: bytes) -> list[Piece]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    pieces = []
    for i, slide in enumerate(prs.slides, start=1):
        title = None
        texts = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
            if not text:
                continue
            if title is None and shape == getattr(slide.shapes, "title", None):
                title = text[:200]
            else:
                texts.append(text)
        if texts or title:
            pieces.append(Piece(title, i, "\n\n".join(texts) or (title or "")))
    return pieces


def _xlsx(data: bytes) -> list[Piece]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pieces = []
    for ws in wb.worksheets[:20]:
        lines = []
        for row in ws.iter_rows(values_only=True, max_row=5000):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                lines.append(" | ".join(cells).rstrip(" |"))
        if lines:
            pieces.append(Piece(ws.title, None, "\n".join(lines)))
    wb.close()
    return pieces


def _csv(data: bytes) -> list[Piece]:
    text = data.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    return [Piece(None, None, "\n".join(" | ".join(r) for r in rows[:5000]))]


def extract(filename: str, mime: str, data: bytes) -> dict:
    name = filename.lower()
    if mime == "application/pdf" or name.endswith(".pdf"):
        pieces = _pdf(data)
    elif "wordprocessingml" in mime or name.endswith(".docx"):
        pieces = _docx(data)
    elif "presentationml" in mime or name.endswith(".pptx"):
        pieces = _pptx(data)
    elif "spreadsheetml" in mime or name.endswith(".xlsx"):
        pieces = _xlsx(data)
    elif mime == "text/csv" or name.endswith(".csv"):
        pieces = _csv(data)
    elif mime.startswith("text/") or name.endswith((".txt", ".md")):
        pieces = [Piece(None, None, data.decode("utf-8-sig", errors="replace"))]
    elif name.endswith((".doc", ".xls", ".ppt")):
        raise ExtractError("legacy_office_format_unsupported")
    else:
        raise ExtractError("unsupported_file_type")
    full = "\n\n".join((f"{p.heading}\n\n" if p.heading else "") + p.text for p in pieces)[:MAX_CHARS]
    return {"text": full, "chunks": _chunk(pieces)}
