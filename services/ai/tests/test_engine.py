import asyncio
import io

import pytest

from qubedu_ai import engine
from qubedu_ai.essay import grade_essay
from qubedu_ai.extract import ExtractError, extract
from qubedu_ai.generate import generate_questions, summarize
from qubedu_ai.mathsolver import option_matches, solve_math
from qubedu_ai.retrieval import Chunk, rank
from qubedu_ai.text import normalize, tokens


def run(coro):
    return asyncio.run(coro)


def test_normalize_matches_typescript_rules():
    assert normalize("  أَحْمَد  ") == "احمد"
    assert normalize("مدرسة") == normalize("مدرسه")
    assert normalize("Café ٣") == "cafe 3"


def test_arithmetic_and_arabic_digits():
    assert solve_math("ما ناتج ٢ + ٣ × ٤ ؟").answer == "14"
    assert solve_math("(1/3) + (1/6)").answer.startswith("1/2")
    assert solve_math("20% of 150").answer == "30"
    assert solve_math("٢٥٪ من ٨٠").answer == "20"


def test_equations():
    lin = solve_math("حل المعادلة 2س + 3 = 7")
    assert lin is not None and lin.answer == "x = 2"
    assert any("2·x = 4" in s for s in lin.steps)
    quad = solve_math("solve x^2 - 5x + 6 = 0")
    assert quad is not None and quad.answer == "x = 2, x = 3"


def test_calculus_and_algebra():
    assert solve_math("derivative of x^3 + 2x").answer == "3*x^2 + 2"
    assert solve_math("integral of 2x").answer == "x^2 + C"
    assert solve_math("factor x^2 - 9").answer == "(x - 3)*(x + 3)"


def test_unsafe_or_non_math_is_rejected():
    assert solve_math("__import__('os').system('ls')") is None
    assert solve_math("ما هي عاصمة مصر") is None
    assert solve_math("2 ** 99999999") is None
    assert solve_math("اشرح لي السؤال رقم 3") is None


def test_mcq_numeric_match():
    m = solve_math("12 / 4 + 1")
    assert option_matches("4", m.value)
    assert not option_matches("3", m.value)


CHUNKS = [
    {
        "id": "1",
        "title": "Biology ch.1",
        "heading": "The cell",
        "page": 3,
        "text": "The mitochondria is the powerhouse of the cell. It produces ATP, the energy currency of the organism.",
    },
    {"id": "2", "title": "Biology ch.1", "heading": "Nucleus", "page": 4, "text": "The nucleus contains DNA and controls the activities of the cell."},
    {"id": "3", "title": "تاريخ", "heading": "مصر", "page": 1, "text": "تأسست القاهرة عام 969 ميلادية على يد الفاطميين وأصبحت عاصمة مصر."},
]


def test_retrieval_ranks_relevant_chunk():
    hits = rank("What does the mitochondria produce?", [Chunk(**{k: v for k, v in c.items()}) for c in CHUNKS])
    assert hits[0].chunk.id == "1"


def test_solve_retrieval_with_sources_and_arabic():
    res = run(engine.solve({"question": "What does the mitochondria produce?", "chunks": CHUNKS, "language": "en"}, None))
    assert res["kind"] == "retrieval"
    assert "ATP" in res["answer"]
    assert res["sources"][0]["page"] == 3
    ar = run(engine.solve({"question": "متى تأسست القاهرة؟", "chunks": CHUNKS, "language": "ar"}, None))
    assert "969" in ar["answer"]


def test_solve_mcq_from_material():
    res = run(
        engine.solve(
            {"question": "Which organelle controls the cell?", "options": ["Mitochondria", "Nucleus", "Ribosome"], "chunks": CHUNKS, "language": "en"}, None
        )
    )
    assert res["kind"] == "mcq" and res["choice"] == 1


def test_solve_uncertain_without_material():
    res = run(engine.solve({"question": "Who won the 1998 world cup?", "chunks": [], "language": "en"}, None))
    assert res["confidence"] == "uncertain"


def test_synonyms_expand_matching():
    chunks = [{"id": "1", "title": "Econ", "text": "Inflation is a general rise in prices over time."}]
    res = run(engine.solve({"question": "ما هو التضخم؟", "chunks": chunks, "synonyms": {"التضخم": ["inflation"]}, "language": "ar"}, None))
    assert res["kind"] == "retrieval"


def test_essay_grading():
    g = grade_essay(
        "Photosynthesis converts light energy into chemical energy. Plants release oxygen.",
        "Plants use light, water and carbon dioxide to make glucose and release oxygen.",
        ["converts light energy into chemical energy", "releases oxygen", "uses carbon dioxide"],
        {},
        "en",
    )
    assert [c["matched"] for c in g["coverage"]] == [True, True, False]
    assert 40 <= g["score"] <= 80


def test_summary_and_quiz_generation():
    text = (
        "Photosynthesis takes place in the chloroplast of plant cells. "
        "The chloroplast contains chlorophyll, which absorbs sunlight. "
        "During photosynthesis plants convert carbon dioxide and water into glucose. "
        "Oxygen is released into the atmosphere as a byproduct of photosynthesis. "
        "Glucose provides energy that the plant uses for growth and respiration. "
        "Respiration releases the energy stored in glucose molecules for the cell."
    )
    assert len(summarize(text, 3)) == 3
    qs = generate_questions(text, 3)
    assert qs, "expected generated questions"
    for q in qs:
        assert "_____" in q["prompt"]
        assert len(q["options"]) == 4
        assert 0 <= q["answer"] < 4


def test_assistant_paths():
    kb = [{"question": "How do I reset my password?", "answer": "Settings → Security → Reset password.", "link": "/settings?tab=security"}]
    r = run(engine.assistant({"message": "how can i reset my password", "knowledge": kb, "language": "en"}, None))
    assert "Security" in r["content"] and r["actions"][0]["link"] == "/settings?tab=security"
    g = run(engine.assistant({"message": "مرحبا", "userName": "سارة", "language": "ar"}, None))
    assert "سارة" in g["content"]
    m = run(engine.assistant({"message": "حل المعادلة 3س = 12", "language": "ar"}, None))
    assert "x = 4" in m["content"]
    nav = run(engine.assistant({"message": "أين أجد الاختبارات؟", "language": "ar"}, None))
    assert any(a["link"] == "/quiz" for a in nav["actions"])


def test_extract_docx_pptx_xlsx_and_text():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    d = Document()
    d.add_heading("Chapter One", level=1)
    d.add_paragraph("Water boils at one hundred degrees Celsius at sea level.")
    buf = io.BytesIO()
    d.save(buf)
    res = extract("a.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", buf.getvalue())
    assert res["chunks"][0]["heading"] == "Chapter One"

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Atoms"
    slide.placeholders[1].text = "Atoms are made of protons, neutrons and electrons."
    pbuf = io.BytesIO()
    prs.save(pbuf)
    pres = extract("s.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", pbuf.getvalue())
    assert pres["chunks"][0]["heading"] == "Atoms" and pres["chunks"][0]["page"] == 1

    wb = Workbook()
    wb.active.append(["Term", "Meaning"])
    wb.active.append(["Velocity", "Speed with direction"])
    xbuf = io.BytesIO()
    wb.save(xbuf)
    x = extract("t.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", xbuf.getvalue())
    assert "Velocity" in x["text"]

    assert "hello" in extract("n.txt", "text/plain", b"hello world")["text"]
    with pytest.raises(ExtractError):
        extract("old.doc", "application/msword", b"\xd0\xcf\x11\xe0")


def test_tokens_drop_stopwords():
    assert "the" not in tokens("the cell and the nucleus")
